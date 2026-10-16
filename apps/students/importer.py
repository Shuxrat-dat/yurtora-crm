"""Excel (.xlsx) dan o'quvchilarni ommaviy import qilish (masalan, Soff CRM'dan o'tish uchun).

Ish tartibi: fayl to'liq tekshiriladi; birorta qatorda xato bo'lsa, HECH NARSA saqlanmaydi
(hammasi yoki hech narsa). Allaqachon mavjud (telefon + ism + familiya) yoki fayl ichida
takrorlangan qatorlar o'tkazib yuboriladi va hisobotda ko'rsatiladi.
"""

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.students.models import Student

MAX_FILE_SIZE = 2 * 1024 * 1024  # 2 MB
MAX_ROWS = 2000
MAX_UNPACKED_SIZE = 30 * 1024 * 1024  # ochilgan hajm: kichik arxiv ichida ulkan fayl (zip-bomba) bo'lmasligi uchun
MAX_ZIP_ENTRIES = 200

# Ustun sarlavhalari (taqqoslashda kichik harf, apostrof variantlari normallashtiriladi)
COLUMNS = [
    "Familiya",
    "Ism",
    "Otasining ismi",
    "Telefon",
    "Tug'ilgan sana",
    "Jinsi",
    "Maktab",
    "Manba",
    "Manzil",
]
REQUIRED = {"familiya", "ism", "telefon"}

GENDERS = {"erkak": "male", "e": "male", "m": "male", "ayol": "female", "a": "female", "f": "female"}
SOURCES = {label.lower(): value for value, label in Student.Source.choices}
SOURCES.update({value: value for value in Student.Source.values})


@dataclass
class ImportResult:
    rows: list = field(default_factory=list)
    errors: list = field(default_factory=list)  # [(qator_raqami, xabar)]
    skipped: list = field(default_factory=list)  # [(qator_raqami, xabar)]
    created: int = 0

    @property
    def ok(self):
        return not self.errors


def _norm_header(value):
    return re.sub(r"[`‘’ʻʼ']", "'", str(value or "")).strip().lower()


def normalize_phone(raw):
    """+998XXXXXXXXX ko'rinishiga keltiradi; yaroqsiz bo'lsa None."""
    digits = re.sub(r"\D", "", str(raw or ""))
    if len(digits) == 9:
        digits = "998" + digits
    if len(digits) == 12 and digits.startswith("998"):
        return "+" + digits
    return None


def _parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"sana noto'g'ri: «{text}» (kutilgan format: 31.12.2010)")


def _cell(row, index):
    if index is None or index >= len(row):
        return ""
    value = row[index]
    return "" if value is None else value


def parse_students(file_obj):
    """Faylni o'qib tekshiradi. Bazaga hech narsa yozmaydi."""
    result = ImportResult()

    data = file_obj.read()
    if len(data) > MAX_FILE_SIZE:
        result.errors.append((0, "Fayl hajmi 2 MB dan oshmasligi kerak."))
        return result
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES or sum(e.file_size for e in entries) > MAX_UNPACKED_SIZE:
                result.errors.append((0, "Fayl juda katta yoki shubhali tuzilgan (ochilgan hajmi chegaradan oshdi)."))
                return result
    except BadZipFile:
        result.errors.append((0, "Faylni o'qib bo'lmadi. .xlsx formatidagi fayl yuklang."))
        return result
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError):
        result.errors.append((0, "Faylni o'qib bo'lmadi. .xlsx formatidagi fayl yuklang."))
        return result

    if not rows:
        result.errors.append((0, "Fayl bo'sh."))
        return result
    if len(rows) - 1 > MAX_ROWS:
        result.errors.append((0, f"Bir martada {MAX_ROWS} tadan ortiq qator import qilib bo'lmaydi."))
        return result

    header = {_norm_header(h): i for i, h in enumerate(rows[0]) if h}
    missing = [c for c in sorted(REQUIRED) if c not in header]
    if missing:
        result.errors.append((1, f"Majburiy ustunlar topilmadi: {', '.join(missing)}."))
        return result

    existing = {
        (s.phone, s.first_name.lower(), s.last_name.lower())
        for s in Student.objects.only("phone", "first_name", "last_name")
    }
    seen = set()

    for number, row in enumerate(rows[1:], start=2):
        if not any(c not in (None, "") for c in row):
            continue  # bo'sh qator

        def get(name, row=row):
            return _cell(row, header.get(name))

        last_name = str(get("familiya")).strip()
        first_name = str(get("ism")).strip()
        problems = []
        if not last_name:
            problems.append("familiya bo'sh")
        if not first_name:
            problems.append("ism bo'sh")
        phone = normalize_phone(get("telefon"))
        if phone is None:
            problems.append(f"telefon noto'g'ri: «{get('telefon')}»")

        birth_date = None
        try:
            birth_date = _parse_date(get("tug'ilgan sana"))
        except ValueError as exc:
            problems.append(str(exc))

        gender_raw = str(get("jinsi")).strip().lower()
        gender = GENDERS.get(gender_raw, "")
        if gender_raw and not gender:
            problems.append(f"jins noto'g'ri: «{gender_raw}» (Erkak/Ayol)")

        source_raw = str(get("manba")).strip().lower()
        source = SOURCES.get(source_raw, "")
        if source_raw and not source:
            problems.append(f"manba noto'g'ri: «{source_raw}»")

        if problems:
            result.errors.append((number, "; ".join(problems)))
            continue

        key = (phone, first_name.lower(), last_name.lower())
        if key in existing:
            result.skipped.append((number, f"{last_name} {first_name}: bazada allaqachon bor"))
            continue
        if key in seen:
            result.skipped.append((number, f"{last_name} {first_name}: faylda takrorlangan"))
            continue
        seen.add(key)

        result.rows.append(
            {
                "last_name": last_name,
                "first_name": first_name,
                "middle_name": str(get("otasining ismi")).strip(),
                "phone": phone,
                "birth_date": birth_date,
                "gender": gender,
                "school": str(get("maktab")).strip(),
                "source": source,
                "address": str(get("manzil")).strip(),
            }
        )
    return result


def import_students(file_obj):
    """Tekshiradi va xato bo'lmasa bitta tranzaksiyada saqlaydi."""
    result = parse_students(file_obj)
    if result.ok and result.rows:
        with transaction.atomic():
            Student.objects.bulk_create([Student(**row) for row in result.rows])
        result.created = len(result.rows)
    return result


def build_template():
    """Bo'sh import shabloni (.xlsx) baytlarini qaytaradi."""
    wb = Workbook()
    ws = wb.active
    ws.title = "O'quvchilar"
    ws.append(COLUMNS)
    ws.append(
        ["Valiyev", "Ali", "Karimovich", "+998901234567", "15.03.2012", "Erkak", "12-maktab", "Instagram", "Toshkent"]
    )
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
