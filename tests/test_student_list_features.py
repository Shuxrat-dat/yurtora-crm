"""O'quvchilar ro'yxati: kurs va to'lov filtrlari, balans, guruh ma'lumoti, Excel eksport."""

from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.courses.models import Course
from apps.finance.services import add_adjustment
from apps.groups.models import Enrollment, Group
from apps.students.models import Student
from apps.teachers.models import Teacher

pytestmark = pytest.mark.django_db
URL = "students:student_list"


@pytest.fixture
def world(director_user):
    eng = Course.objects.create(name="English", monthly_price=400000)
    math = Course.objects.create(name="Math", monthly_price=300000)
    teacher = Teacher.objects.create(full_name="Ustoz", phone="+998900000000")
    g_eng = Group.objects.create(name="ENG-1", course=eng, teacher=teacher, start_date="2026-01-01")
    g_math = Group.objects.create(name="MATH-1", course=math, teacher=teacher, start_date="2026-01-01")
    debtor = Student.objects.create(first_name="Qarzdor", last_name="Aliyev", phone="+998901110001")
    clean = Student.objects.create(first_name="Toza", last_name="Valiyev", phone="+998901110002")
    for g in (g_eng, g_math):
        Enrollment.objects.create(student=debtor, group=g, status="active", start_date="2026-01-01", paid_from="2026-01-01")
    Enrollment.objects.create(student=clean, group=g_math, status="active", start_date="2026-01-01", paid_from="2026-01-01")
    add_adjustment(debtor, g_eng, -300000, director_user, "qarz")
    add_adjustment(debtor, g_math, -100000, director_user, "qarz")
    add_adjustment(clean, g_math, 50000, director_user, "avans")
    return {"eng": eng, "math": math, "g_eng": g_eng, "debtor": debtor, "clean": clean}


def names(resp):
    return sorted(s.first_name for s in resp.context["students"])


def test_payment_filter(client, admin_user, world):
    client.force_login(admin_user)
    assert names(client.get(reverse(URL), {"payment": "debt"})) == ["Qarzdor"]
    assert names(client.get(reverse(URL), {"payment": "ok"})) == ["Toza"]


def test_course_filter(client, admin_user, world):
    client.force_login(admin_user)
    assert names(client.get(reverse(URL), {"course": world["eng"].pk})) == ["Qarzdor"]
    assert names(client.get(reverse(URL), {"course": world["math"].pk})) == ["Qarzdor", "Toza"]


def test_balance_not_inflated_by_joins(client, admin_user, world):
    """Ikki guruhli o'quvchi guruh/kurs filtri bilan ham to'g'ri umumiy balansga ega."""
    client.force_login(admin_user)
    resp = client.get(reverse(URL), {"course": world["math"].pk})
    debtor = next(s for s in resp.context["students"] if s.first_name == "Qarzdor")
    assert debtor.balance == -400000
    per_group = {e.group.name: e.group_balance for e in debtor.open_enrollments}
    assert per_group == {"ENG-1": -300000, "MATH-1": -100000}


def test_teacher_does_not_see_balances_or_export(client, teacher_user, world):
    client.force_login(teacher_user)
    assert client.get(reverse(URL), {"export": "xlsx"}).status_code == 403


def test_excel_export_respects_filters(client, admin_user, world):
    client.force_login(admin_user)
    resp = client.get(reverse(URL), {"payment": "debt", "export": "xlsx"})
    assert resp.status_code == 200
    assert resp["Content-Disposition"].startswith("attachment")
    ws = load_workbook(BytesIO(resp.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][1] == "Familiya"
    assert len(rows) == 2 and rows[1][2] == "Qarzdor" and rows[1][8] == -400000


def test_row_menu_has_quick_payment(client, admin_user, world):
    client.force_login(admin_user)
    html = client.get(reverse(URL)).content.decode()
    assert reverse("finance:payment_modal_form", args=[world["debtor"].pk]) in html
    assert "Topildi:" in html
